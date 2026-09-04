"""将 Campaign 行动 + 画像 + 反馈 + 判定结果打包为 Excel。"""
from __future__ import annotations

from datetime import datetime
from io import BytesIO
from typing import Any, Dict, List, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from schemas import Persona, SimulateResult
from services.job_store import load_outcomes
from services.simulate import INTERVENTION_LABELS, intervention_label

DECISION_CN = {
    "advance": "推进",
    "hesitate": "犹豫",
    "reject": "拒绝",
    "na": "无关",
}

HEADER_FILL = PatternFill("solid", fgColor="1B3A4B")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
WRAP = Alignment(wrap_text=True, vertical="top")


def _outcome_label_map() -> Dict[str, str]:
    out: Dict[str, str] = {}
    for o in load_outcomes():
        out[o.id] = o.label or o.name or o.id
    return out


def _fmt_outcomes(ids: List[str], labels: Dict[str, str]) -> str:
    if not ids:
        return ""
    return "、".join(labels.get(i, i) for i in ids)


def _style_header(ws, ncols: int) -> None:
    for col in range(1, ncols + 1):
        cell = ws.cell(1, col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)


def _autosize(ws, max_width: int = 48) -> None:
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        width = 10
        for cell in col:
            val = "" if cell.value is None else str(cell.value)
            width = max(width, min(max_width, len(val) + 2))
        ws.column_dimensions[letter].width = width


def _append_row(ws, values: List[Any]) -> None:
    ws.append(values)
    for cell in ws[ws.max_row]:
        cell.alignment = WRAP


def build_campaign_workbook(
    *,
    campaign: str,
    results: List[SimulateResult],
    personas: List[Persona],
    campaign_hits: Optional[List[str]] = None,
    interventions: Optional[List[str]] = None,
    use_llm: bool = False,
    use_memory: bool = False,
) -> bytes:
    """生成多工作表 xlsx 字节流。"""
    personas_by_id = {p.id: p for p in personas}
    labels = _outcome_label_map()
    hits = campaign_hits or []
    ints = interventions or []
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    summary = {"advance": 0, "hesitate": 0, "reject": 0, "na": 0}
    for r in results:
        if r.decision in summary:
            summary[r.decision] += 1

    wb = Workbook()

    # —— 1. 行动 ——
    ws_act = wb.active
    ws_act.title = "行动"
    ws_act.append(["字段", "内容"])
    _style_header(ws_act, 2)
    act_rows = [
        ("导出时间", now),
        ("测试模式", "独立思考（LLM）" if use_llm else "规则引擎"),
        ("消费者记忆", "开启" if use_memory else "关闭"),
        ("测试人数", len(results)),
        ("推进", summary["advance"]),
        ("犹豫", summary["hesitate"]),
        ("拒绝", summary["reject"]),
        ("无关", summary["na"]),
        ("识别维度", "、".join(hits) if hits else "—"),
        (
            "识别助力点",
            "、".join(intervention_label(i) if i in INTERVENTION_LABELS else i for i in ints)
            if ints
            else "—",
        ),
        ("行动/话术全文", (campaign or "").strip() or "—"),
    ]
    for k, v in act_rows:
        _append_row(ws_act, [k, v])
    ws_act.column_dimensions["A"].width = 16
    ws_act.column_dimensions["B"].width = 80

    # —— 2. 结果总览 ——
    ws_res = wb.create_sheet("结果总览")
    res_headers = [
        "消费者ID",
        "姓名",
        "分群",
        "判定",
        "意愿分",
        "下一步",
        "已改善",
        "未解决",
        "模式",
        "真实验证",
        "验证备注",
    ]
    ws_res.append(res_headers)
    _style_header(ws_res, len(res_headers))
    for r in results:
        p = personas_by_id.get(r.persona_id)
        _append_row(
            ws_res,
            [
                r.persona_id,
                r.persona_name,
                p.segment if p else "",
                DECISION_CN.get(r.decision, r.decision),
                r.willingness if r.willingness is not None else "",
                r.next_step or "",
                _fmt_outcomes(r.outcome_improved, labels),
                _fmt_outcomes(r.unresolved_outcomes, labels),
                "LLM" if r.mode == "llm" else "规则",
                r.verification or "",
                r.verification_note or "",
            ],
        )
    _autosize(ws_res)

    # —— 3. 反馈 ——
    ws_fb = wb.create_sheet("反馈")
    fb_headers = [
        "消费者ID",
        "姓名",
        "判定",
        "意愿分",
        "第一人称反馈",
        "内心推理",
        "命中维度",
        "激活干预",
    ]
    ws_fb.append(fb_headers)
    _style_header(ws_fb, len(fb_headers))
    for r in results:
        _append_row(
            ws_fb,
            [
                r.persona_id,
                r.persona_name,
                DECISION_CN.get(r.decision, r.decision),
                r.willingness if r.willingness is not None else "",
                r.reaction or "",
                r.reasoning or "",
                "、".join(r.hit_factors or r.activated_factors or []),
                "、".join(
                    intervention_label(i) if i in INTERVENTION_LABELS else i
                    for i in (r.interventions or [])
                ),
            ],
        )
    _autosize(ws_fb, max_width=60)

    # —— 4. 画像 ——
    ws_p = wb.create_sheet("画像")
    p_headers = [
        "消费者ID",
        "姓名",
        "年龄",
        "性别",
        "城市",
        "职业",
        "收入",
        "家庭",
        "分群",
        "角色",
        "对象",
        "决策风格",
        "OSA阶段",
        "OSA严重度",
        "起点情境",
        "Job",
        "当前卡点",
        "任务陈述",
        "核心动机",
        "恐惧",
        "决策逻辑",
        "口头禅",
        "阻碍",
        "驱动",
        "主导维度",
        "期望结果(在乎/满意)",
    ]
    ws_p.append(p_headers)
    _style_header(ws_p, len(p_headers))
    seen = set()
    for r in results:
        if r.persona_id in seen:
            continue
        seen.add(r.persona_id)
        p = personas_by_id.get(r.persona_id)
        if not p:
            _append_row(ws_p, [r.persona_id, r.persona_name] + [""] * (len(p_headers) - 2))
            continue
        j = p.jtbd
        ms = p.mindset
        dos = "；".join(
            f"{labels.get(d.id, d.id)}({d.importance}/{d.satisfaction})"
            for d in (j.desired_outcomes or [])
        )
        dom = "、".join(f"{d.code}:{d.weight}" for d in (p.dominant_features or []))
        _append_row(
            ws_p,
            [
                p.id,
                p.name,
                p.age,
                p.gender,
                p.city,
                p.occupation,
                p.income,
                p.family,
                p.segment,
                p.role,
                p.subject,
                p.decision_style,
                p.osa.stage if p.osa else "",
                p.osa.severity if p.osa else "",
                j.entry_situation or "",
                j.core_job or j.job_id or "",
                j.current_step or "",
                j.job_statement or "",
                ms.core_motive if ms else "",
                ms.fear if ms else "",
                ms.decision_logic if ms else "",
                ms.quote if ms else "",
                "；".join(p.blockers or []),
                "；".join(p.drivers or []),
                dom,
                dos,
            ],
        )
    _autosize(ws_p, max_width=40)

    # —— 5. 关注点矩阵（便于筛选）——
    ws_m = wb.create_sheet("关注点矩阵")
    all_oids: List[str] = []
    for r in results:
        for oid in list(r.outcome_improved or []) + list(r.unresolved_outcomes or []):
            if oid not in all_oids:
                all_oids.append(oid)
    m_headers = ["消费者ID", "姓名", "判定", "意愿分"] + [
        labels.get(oid, oid) for oid in all_oids
    ]
    ws_m.append(m_headers)
    _style_header(ws_m, len(m_headers))
    for r in results:
        improved = set(r.outcome_improved or [])
        unresolved = set(r.unresolved_outcomes or [])
        cells = [
            r.persona_id,
            r.persona_name,
            DECISION_CN.get(r.decision, r.decision),
            r.willingness if r.willingness is not None else "",
        ]
        for oid in all_oids:
            if oid in improved:
                cells.append("改善")
            elif oid in unresolved:
                cells.append("未解")
            else:
                cells.append("—")
        _append_row(ws_m, cells)
    _autosize(ws_m)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
